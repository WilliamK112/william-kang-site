from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import shutil

ROOT=Path(__file__).resolve().parent.parent
BLUE='4472C4'; BLACK='000000'; LINK='0563C1'

contact=[('Madison, WI',None),('347-866-8326',None),('ckang53@wisc.edu','mailto:ckang53@wisc.edu'),('linkedin.com/in/kcw2027','https://www.linkedin.com/in/kcw2027/'),('github.com/WilliamK112','https://github.com/WilliamK112')]

def set_font(run, name='Arial', size=9.4, bold=None, italic=None, color=BLACK):
    run.font.name=name; run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),name); run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),name)
    run.font.size=Pt(size); run.font.color.rgb=RGBColor.from_string(color)
    if bold is not None: run.bold=bold
    if italic is not None: run.italic=italic

def add_link(p,label,url,size=9):
    rid=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),rid); r=OxmlElement('w:r'); rp=OxmlElement('w:rPr')
    for tag,val in [('w:rFonts',None),('w:color',LINK),('w:u','single'),('w:sz',str(int(size*2))),('w:szCs',str(int(size*2)))]:
        el=OxmlElement(tag)
        if tag=='w:rFonts': el.set(qn('w:ascii'),'Arial'); el.set(qn('w:hAnsi'),'Arial')
        else: el.set(qn('w:val'),val)
        rp.append(el)
    r.append(rp); t=OxmlElement('w:t'); t.text=label; r.append(t); h.append(r); p._p.append(h)

def base_doc(master=False):
    d=Document(); sec=d.sections[0]
    sec.top_margin=Inches(.38 if not master else .46); sec.bottom_margin=Inches(.38 if not master else .46)
    sec.left_margin=Inches(.48 if not master else .55); sec.right_margin=Inches(.48 if not master else .55)
    sec.header_distance=Inches(.2); sec.footer_distance=Inches(.2)
    normal=d.styles['Normal']; normal.font.name='Arial'; normal._element.rPr.rFonts.set(qn('w:ascii'),'Arial'); normal._element.rPr.rFonts.set(qn('w:hAnsi'),'Arial'); normal.font.size=Pt(9.4 if not master else 9.2)
    normal.paragraph_format.space_after=Pt(0); normal.paragraph_format.line_spacing=1
    title=d.add_paragraph(); title.alignment=WD_ALIGN_PARAGRAPH.CENTER; title.paragraph_format.space_after=Pt(1)
    set_font(title.add_run('CHING-WEI (WILLIAM) KANG'),size=15,bold=True)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(1.5)
    for i,(label,url) in enumerate(contact):
        if i: set_font(p.add_run(' | '),size=9)
        if url: add_link(p,label,url,9)
        else: set_font(p.add_run(label),size=9)
    return d

def bottom_border(p):
    pPr=p._p.get_or_add_pPr(); pbdr=pPr.find(qn('w:pBdr'))
    if pbdr is None: pbdr=OxmlElement('w:pBdr'); pPr.append(pbdr)
    b=OxmlElement('w:bottom'); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'10'); b.set(qn('w:space'),'1'); b.set(qn('w:color'),BLACK); pbdr.append(b)

def section(d,text,master=False):
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(3 if not master else 4); p.paragraph_format.space_after=Pt(1); p.paragraph_format.keep_with_next=True
    set_font(p.add_run(text.upper()),size=10.6,bold=True,color=BLUE); bottom_border(p); return p

def lr(d,left,right,bold_left=True,italic_right=False,size=9.5,before=0,after=0):
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after); p.paragraph_format.keep_with_next=True
    usable=d.sections[0].page_width-d.sections[0].left_margin-d.sections[0].right_margin
    p.paragraph_format.tab_stops.add_tab_stop(usable,WD_TAB_ALIGNMENT.RIGHT)
    set_font(p.add_run(left),size=size,bold=bold_left); p.add_run('\t'); set_font(p.add_run(right),size=size,bold=bold_left if not italic_right else False,italic=italic_right)
    return p

def role(d,title,date,size=9.35):
    return lr(d,title,date,bold_left=True,italic_right=True,size=size)

def bullet(d,text,lead=None,size=9.35,master=False):
    p=d.add_paragraph(style=None); pf=p.paragraph_format; pf.left_indent=Inches(.18); pf.first_line_indent=Inches(-.12); pf.space_after=Pt(.55 if not master else .8); pf.line_spacing=1.0; pf.keep_together=True
    set_font(p.add_run('•  '),size=size,bold=True)
    if lead:
        effective_lead=lead
        if not text.startswith(effective_lead):
            markers=(' by ',' through ',', using ',', integrating ',', then ',' with ')
            positions=[text.find(marker) for marker in markers if text.find(marker)>0]
            effective_lead=text[:min(positions)] if positions else text.split(',',1)[0]
        set_font(p.add_run(effective_lead),size=size,bold=True); set_font(p.add_run(text[len(effective_lead):]),size=size)
    else: set_font(p.add_run(text),size=size)
    return p

def textline(d,label,value,size=9.35):
    p=d.add_paragraph(); p.paragraph_format.space_after=Pt(.45)
    set_font(p.add_run(label),size=size,bold=True); set_font(p.add_run(value),size=size)

def education(d,include_gwu=True,master=False,coursework=None):
    section(d,'Education',master)
    lr(d,'University of Wisconsin-Madison','Madison, WI',size=9.6)
    lr(d,'B.S. Computer Science and Data Science | GPA: 3.91/4.0','Expected May 2027',bold_left=False,italic_right=True,size=9.25)
    if coursework: textline(d,'Relevant coursework: ',coursework,size=9.15)
    if include_gwu:
        lr(d,'The George Washington University','Washington, DC',size=9.6,before=.5)
        lr(d,'Undergraduate coursework in Computer Science and Data Science | GPA: 3.94/4.0','2023 - 2024',bold_left=False,italic_right=True,size=9.15)

def experience(d,company,location,title,date,bullets,master=False):
    lr(d,company,location,size=9.55,before=.6)
    role(d,title,date,size=9.25)
    for t,lead in bullets: bullet(d,t,lead,size=9.25 if not master else 9.1,master=master)

def project(d,name,date,tech,bullets,master=False):
    lr(d,name,date,size=9.45,before=.45)
    textline(d,'Technologies: ',tech,size=9.05)
    for t,lead in bullets: bullet(d,t,lead,size=9.2 if not master else 9.05,master=master)
    return d.paragraphs[-(2+len(bullets))]

def achievements(d,items,master=False):
    section(d,'Achievements',master)
    for item in items: bullet(d,item,None,size=9.15 if not master else 9.05,master=master)

def save(d,subdir,filename):
    out=ROOT/'resumes'/subdir; out.mkdir(parents=True,exist_ok=True); path=out/filename; d.save(path); return path

# SWE
swe=base_doc(); education(swe,include_gwu=False,coursework='Data Structures & Algorithms, Artificial Intelligence, Computer Engineering, Discrete Math, Statistical Modeling, Linear Algebra.')
section(swe,'Skills'); textline(swe,'Programming: ','Python, TypeScript, Java, SQL, Go, C++'); textline(swe,'Technologies: ','Next.js, React, Node.js, FastAPI, Express, PostgreSQL, Redis, SQLite, Docker, AWS, REST APIs, RBAC, SSE, CI/CD')
section(swe,'Professional Experience')
experience(swe,'Global AI (Global API Inc.)','New York, NY | Hybrid','Full-Stack Engineering Intern','May 2026 - Present',[
('Led backend development to make an AI video generation platform production-ready, using FastAPI, Gemini, PostgreSQL, and Redis to orchestrate validated multistage jobs with asynchronous execution and failure recovery.','Led backend development'),
('Reduced redundant model work and release risk by implementing parallel execution, cache-backed processing, RBAC, observability, and CI/CD, improving reliability across long-running generation workflows and repeatable production releases.','Implemented parallel task execution'),
('Aligned engineering delivery with company priorities by partnering with the CEO and four departments, translating product requirements into an SSE-streamed dashboard that improved real-time review, debugging, and release visibility.','Partnered directly with the CEO and four departments')])
experience(swe,'Guangdong Akang Health Technology Group Co., Ltd.','China | On-site','AI & Software Engineering Intern','Jul 2026 - Aug 2026',[
('Made fragmented pharmaceutical evidence searchable and reusable by building interfaces and JSON/SQLite pipelines that normalized reports, legal evidence, company, and partner records with provenance, validation metadata, and reproducible exports.','Built searchable research interfaces and structured JSON/SQLite pipelines'),
('Improved trust in AI-assisted research by combining automated completeness, citation, and source-coverage checks with human review checkpoints, producing auditable evidence for downstream pharmaceutical decisions.','Implemented automated checks')])
experience(swe,'China Association for Science and Technology (CAST)','Washington, DC | Remote','Research Assistant','Dec 2024 - Feb 2025',[
('Supported national research decisions on AI talent allocation by reconciling policy and institutional sources in Python, building a documented linear forecasting model, and delivering reproducible evidence to project stakeholders.','Built a Python linear model')])
experience(swe,'The Marcus Harris Foundation','Washington, DC','Data Engineering and Automation Intern','Sep 2024 - Dec 2024',[
('Made fragmented nonprofit registry data usable for reporting by building a Python ingestion service with BeautifulSoup, Requests, and Pandas, adding reusable parsers, schema checks, and auditable source fields.','Built a Python ingestion service'),
('Improved downstream reporting quality by loading 50,000+ records into PostgreSQL with 98% accuracy and applying validation and filters across EIN, filing year, tax category, and geography.','Loaded 50,000+ records into PostgreSQL')])
experience(swe,'Springer Capital','Chicago, IL','Data Analyst and Automation Intern','Aug 2024 - Oct 2024',[
('Automated recurring management workflows by integrating Power Automate, Power Apps, Microsoft Lists, and Scoro CRM, reducing fragmented handoffs and improving operational visibility.','Automated recurring management workflows')])
section(swe,'Selected Projects')
project(swe,'ClearMove - Rental Evidence Platform','Oct 2025 - Present','Next.js, TypeScript, PostgreSQL, Drizzle ORM, REST APIs, pdf-lib',[('Shipped a deployed full-stack platform that converts lease and inspection inputs into auditable, timestamped PDFs across 21 API routes and 10 PostgreSQL tables, informed by 50+ stakeholder interviews and production-ready data exports.','Shipped a deployed full-stack platform')])
project(swe,'Multi-Agent OpenClaw','Jan 2026 - Present','TypeScript, Node.js, Express, Zod, REST APIs, persistent run state, automated tests',[('Built an orchestration API with schema validation, persistent run memory, task routing, and QA-gated document exports; covered core workflows with seven automated tests and repeatable QA regression checks across document workflows.','Built an orchestration API')])
project(swe,'TableUs - Group Restaurant Discovery','2026','Next.js, React, Python, FastAPI, Gemini API, Google Maps API',[('Built a full-stack recommendation workflow that combines group preferences, geospatial context, and natural-language input; earned Second Place at the 2026 Cursor Hackathon through an explainable, location-aware group decision flow.','Built a full-stack recommendation workflow')])
project(swe,'Badminton AI Match Analysis','2026','Python, computer vision, player tracking, pose estimation, video processing',[('Developed a two-player analysis pipeline with stable tracking, pose overlays, shuttle recovery, and annotated video outputs for repeatable match review and model debugging.','Developed a two-player analysis pipeline')])
section(swe,'Open Source')
bullet(swe,'Contributed 201 merged pull requests to external repositories, working through maintainer review, automated tests, and CI workflows across production developer tools.','Contributed 201 merged pull requests')
bullet(swe,'Shipped reviewed contributions to GitHub Docs, Docker CLI, Microsoft APM, NVIDIA NemoClaw, Meta DotSlash, Repomix, and additional open-source projects.','Shipped reviewed contributions')
achievements(swe,['Third Prize, 2025 Badger Build Fest Venture Track - ClearMove','Second Place (Skills) and Third Place (Qualification), 2023 VEX Robotics World Championship'])
swe_path=save(swe,'swe','Ching-Wei-Kang-SWE-Resume.docx')

# AI
ai=base_doc(); education(ai,include_gwu=True,coursework='Artificial Intelligence, Statistical Modeling, Linear Algebra, Data Structures & Algorithms, Data Science Programming, Discrete Math.')
section(ai,'AI and Engineering Skills'); textline(ai,'Programming: ','Python, TypeScript, SQL, Java, Go, C++'); textline(ai,'AI systems: ','Gemini API, multi-agent orchestration, computer vision, reinforcement learning, JSON Schema, evaluation, human review'); textline(ai,'Engineering: ','FastAPI, React, Next.js, Node.js, PostgreSQL, Redis, REST APIs, SSE, Docker, CI/CD')
section(ai,'Relevant Experience')
experience(ai,'Global AI (Global API Inc.)','New York, NY | Hybrid','Full-Stack Engineering Intern','May 2026 - Present',[
('Led backend development to make an AI video generation platform production-ready, integrating Gemini inference with FastAPI, PostgreSQL, and Redis for validated multistage workflows with failure recovery.','Led backend development'),
('Reduced repeated model calls and improved inference visibility by orchestrating parallel keyframe and video tasks with cached intermediates, SSE progress streaming, cache observability, and production debugging controls.','Orchestrated parallel keyframe and video tasks'),
('Aligned AI delivery with product and operational needs by partnering with the CEO and four departments, then implementing RBAC, validation, observability, and CI/CD controls for repeatable releases and stakeholder review.','Partnered with the CEO and four departments')])
experience(ai,'Guangdong Akang Health Technology Group Co., Ltd.','China | On-site','AI & Software Engineering Intern','Jul 2026 - Aug 2026',[
('Made pharmaceutical market research traceable by designing multi-agent workflows with evidence grading, human review checkpoints, source provenance, citation checks, and automated validation, enabling reproducible exports.','Designed multi-agent research workflows'),
('Enabled reviewers to verify AI-assisted findings by developing searchable interfaces and structured exports with source IDs, retrieval dates, and citation checks across legal, company, and partner-verification evidence.','Developed searchable interfaces and structured exports')])
experience(ai,'China Association for Science and Technology (CAST)','Washington, DC | Remote','Research Assistant','Dec 2024 - Feb 2025',[
('Supported national research decisions on AI talent allocation by reconciling institutional evidence in Python, documenting model assumptions, and delivering a reproducible forecasting workflow to project stakeholders.','Built a Python forecasting model')])
experience(ai,'The Marcus Harris Foundation','Washington, DC','Data Engineering and Automation Intern','Sep 2024 - Dec 2024',[
('Made public nonprofit records analysis-ready by building a Python ingestion pipeline with reusable parsers, schema validation, and auditable source fields, then loading 50,000+ PostgreSQL records with 98% accuracy.','Built a Python ingestion pipeline')])
section(ai,'Selected AI Projects')
project(ai,'Multi-Agent OpenClaw','Jan 2026 - Present','TypeScript, Node.js, Express, Zod, agent orchestration, persistent memory',[('Built a task-routing agent system that selects workflows by complexity, persists run memory, and gates Markdown/DOCX outputs with automated QA across seven validated tests, persistent execution state, and regression-safe document workflows.','Built a task-routing agent system')])
project(ai,'Badminton AI Match Analysis','2026','Python, computer vision, player tracking, pose estimation, video processing',[('Built a computer-vision pipeline that tracks two players, overlays pose estimates, recovers shuttle motion, and renders annotated video for repeatable match analysis, diagnostic review, structured performance comparison, and model debugging.','Built a computer-vision pipeline')])
project(ai,'Adaptive Thinking and Reinforcement Learning','2026','Python, PyTorch, DreamerV3, PPO, adaptive halting, experiment diagnostics',[('Benchmarked adaptive-computation and reinforcement-learning workflows, documenting sparse-retrieval failures, training behavior, reproducible evaluation results, and failure-mode analysis across adaptive computation experiments.','Benchmarked adaptive-computation and reinforcement-learning workflows')])
project(ai,'ClearMove - AI Rental Evidence Platform','Oct 2025 - Present','Gemini API, Next.js, TypeScript, PostgreSQL, Drizzle ORM, pdf-lib',[('Integrated Gemini into a deployed evidence workflow that converts lease and inspection inputs into auditable PDFs, informed by 50+ stakeholder interviews, traceable data sources, production-ready evidence exports, and downstream review workflows.','Integrated Gemini into a deployed evidence workflow')])
project(ai,'TableUs - AI Group Recommendation','2026','Python, FastAPI, Gemini API, Google Maps API, preference modeling',[('Combined group preferences, geospatial context, and natural-language input into explainable recommendations; earned Second Place at the 2026 Cursor Hackathon.','Combined group preferences, geospatial context, and natural-language input')])
section(ai,'Open Source')
bullet(ai,'Contributed 201 merged pull requests to external repositories through maintainer review and CI workflows, including GitHub Docs, Docker CLI, Microsoft APM, NVIDIA NemoClaw, Meta DotSlash, and Repomix.','Contributed 201 merged pull requests')
achievements(ai,['Second Place, 2026 Cursor Hackathon - TableUs','Third Prize, 2025 Badger Build Fest Venture Track - ClearMove'])
ai_path=save(ai,'ai','Ching-Wei-Kang-AI-Resume.docx')

# Data Science
ds=base_doc(); education(ds,include_gwu=True,coursework='Statistical Modeling, Linear Algebra, Data Science Programming, Artificial Intelligence, Data Structures & Algorithms, Discrete Math.')
section(ds,'Data Skills'); textline(ds,'Programming: ','Python, SQL, R, Java, TypeScript'); textline(ds,'Analytics: ','Pandas, Excel, Power BI, Plotly, Matplotlib, statistical modeling, data visualization, forecasting, data quality'); textline(ds,'Data systems: ','PostgreSQL, SQLite, REST APIs, BeautifulSoup, Requests, Power Automate, Power Apps, Microsoft Lists')
section(ds,'Data and Research Experience')
experience(ds,'Global AI (Global API Inc.)','New York, NY | Hybrid','Full-Stack Engineering Intern','May 2026 - Present',[
('Made long-running AI generation jobs reliable and measurable by building FastAPI workflows backed by PostgreSQL and Redis, with validation, cache-backed processing, failure recovery, and operational observability.','Built PostgreSQL and Redis-backed FastAPI workflows'),
('Improved cross-functional visibility by partnering with the CEO and four departments to translate reporting needs into an SSE-streamed dashboard for real-time progress, debugging, and release monitoring.','Partnered with the CEO and four departments')])
experience(ds,'China Association for Science and Technology (CAST)','Washington, DC | Remote','Research Assistant','Dec 2024 - Feb 2025',[
('Supported national AI talent-allocation research by reconciling multi-source policy evidence and building a documented Python linear model, delivering reproducible forecasts for project stakeholders.','Analyzed AI policy and talent-allocation trends')])
experience(ds,'The Marcus Harris Foundation','Washington, DC','Data Engineering and Automation Intern','Sep 2024 - Dec 2024',[
('Made public nonprofit tax records analysis-ready by building a Python pipeline with BeautifulSoup, Requests, and Pandas, using reusable parsers, schema validation, and auditable source fields.','Built a Python ingestion pipeline'),
('Improved reporting quality by loading 50,000+ records into PostgreSQL with 98% accuracy and applying filters across EIN, filing year, tax category, and geography for reproducible management analysis.','Parsed and loaded 50,000+ records into PostgreSQL')])
experience(ds,'Springer Capital','Chicago, IL','Data Analyst and Automation Intern','Aug 2024 - Oct 2024',[
('Enabled recurring management decisions by consolidating operational data from databases, APIs, spreadsheets, and web sources with Python, SQL, Excel, and Power BI into KPI and reporting workflows.','Collected and analyzed operational data'),
('Reduced fragmented reporting handoffs by connecting Power Automate, Power Apps, Microsoft Lists, and Scoro CRM, improving data consistency and operational visibility across recurring processes.','Connected Power Automate, Power Apps, Microsoft Lists, and Scoro CRM workflows')])
experience(ds,'Chihuo Inc.','Remote','Data Analyst','Jun 2024 - Sep 2024',[
('Supported North American market strategy by analyzing demographic, behavioral, and preference data with Python and Excel, contributing visualizations and recommendations to a 200+ page Asian food industry white paper.','Analyzed demographics, customer behavior, and consumer preferences')])
experience(ds,'Guangdong Akang Health Technology Group Co., Ltd.','China | On-site','AI & Software Engineering Intern','Jul 2026 - Aug 2026',[
('Made multi-source pharmaceutical research auditable by normalizing records into JSON and SQLite datasets with source IDs, retrieval dates, validation metadata, citation checks, and reproducible exports.','Normalized multi-source pharmaceutical research records')])
section(ds,'Selected Data Projects')
project(ds,'ClearMove - Rental Evidence Data Platform','Oct 2025 - Present','PostgreSQL, SQL, Drizzle ORM, TypeScript, Gemini API, pdf-lib',[('Designed data flows across 10 PostgreSQL tables and 21 API routes, converting lease and inspection inputs into traceable evidence records and timestamped PDF exports with traceable inputs, validation metadata, and reproducible transformations.','Designed data flows across 10 PostgreSQL tables and 21 API routes')])
project(ds,'TableUs - Group Recommendation System','2026','Python, FastAPI, Gemini API, Google Maps API, preference modeling',[('Combined group preferences, geospatial context, and natural-language input into explainable restaurant recommendations; earned Second Place at the 2026 Cursor Hackathon through an explainable, location-aware group decision flow.','Combined group preferences, geospatial context, and natural-language input')])
project(ds,'Badminton Match Analytics','2026','Python, computer vision, tracking, pose estimation, video visualization',[('Produced structured two-player tracking, pose overlays, shuttle recovery, and annotated video outputs for repeatable performance analysis, structured comparison, and repeatable review of player movement and rally-level outputs.','Produced structured two-player tracking')])
section(ds,'Open Source')
bullet(ds,'Contributed 201 merged pull requests to external repositories through maintainer review and CI workflows, including GitHub Docs, Docker CLI, Microsoft APM, and NVIDIA NemoClaw.','Contributed 201 merged pull requests')
textline(ds,'Leadership: ','Badger Blockchain E-Board, Social Media and Communications (Summer 2026 - Present); Badminton E-Board (Spring 2025 - Present).',9.05)
achievements(ds,['Second Place: 2024 GWU Data Science Association Datathon (Python) and 2026 Cursor Hackathon (TableUs)'])
ds_path=save(ds,'data-science','Ching-Wei-Kang-Data-Science-Resume.docx')

# Master's application CV
m=base_doc(master=True); education(m,include_gwu=True,master=True,coursework='Artificial Intelligence, Data Structures & Algorithms, Statistical Modeling, Linear Algebra, Computer Engineering, Data Science Programming, Discrete Mathematics.')
section(m,'Academic Profile',True)
textline(m,'Interests: ','Applied AI, multi-agent systems, trustworthy data workflows, computer vision, reinforcement learning, and production software systems.',9.1)
textline(m,'Technical foundation: ','Python, TypeScript, Java, SQL, Go, C++, R; PyTorch, FastAPI, React/Next.js, Node.js, PostgreSQL, Redis, Docker, AWS, REST APIs, CI/CD.',9.1)
section(m,'Research and Professional Experience',True)
experience(m,'Global AI (Global API Inc.)','New York, NY | Hybrid','Full-Stack Engineering Intern','May 2026 - Present',[
('Made an AI video generation platform production-ready by leading backend development with Gemini, FastAPI, PostgreSQL, and Redis across validated multistage workflows.','Led backend development'),
('Reduced repeated work and improved operational visibility through parallel execution, caching, RBAC, observability, CI/CD, and SSE streaming, while translating priorities from the CEO and four departments into reliable releases.','Implemented parallel execution, caching, RBAC, observability, CI/CD, and SSE streaming')],True)
experience(m,'Guangdong Akang Health Technology Group Co., Ltd.','China | On-site','AI & Software Engineering Intern','Jul 2026 - Aug 2026',[
('Made pharmaceutical research traceable by designing multi-agent workflows with evidence grading, human review checkpoints, automated validation, and auditable source provenance.','Designed multi-agent pharmaceutical research workflows'),
('Enabled reviewers to verify international healthcare findings by building searchable interfaces and structured JSON/SQLite exports across reports, legal evidence, company data, and partner records.','Built searchable interfaces and structured JSON/SQLite exports')],True)
experience(m,'China Association for Science and Technology (CAST)','Washington, DC | Remote','Research Assistant','Dec 2024 - Feb 2025',[
('Supported national AI talent-allocation research by reconciling institutional evidence and building a documented Python linear model for reproducible forecasting and stakeholder analysis.','Analyzed AI policy and talent-allocation trends')],True)
experience(m,'Springer Capital','Chicago, IL','Data Analyst and Automation Intern','Aug 2024 - Oct 2024',[
('Improved management reporting and process visibility by analyzing operational data with Python, SQL, Excel, and Power BI and connecting Power Platform and CRM workflows.','Analyzed operational data')],True)
experience(m,'Chihuo Inc.','Remote','Data Analyst','Jun 2024 - Sep 2024',[
('Supported North American market strategy by analyzing demographic, behavioral, and preference data with Python and Excel, contributing visualizations and recommendations to a 200+ page Asian food industry white paper.','Analyzed demographic, behavioral, and preference data')],True)
section(m,'Selected Technical and Research Projects',True)
project(m,'Multi-Agent OpenClaw','Jan 2026 - Present','TypeScript, Node.js, Express, Zod, agent orchestration, persistent memory, automated QA',[('Built a multi-agent orchestration system that routes work by task type and complexity, persists run memory, and exports QA-gated Markdown/DOCX artifacts across seven validated tests.','Built a multi-agent orchestration system')],True)
project(m,'Adaptive Thinking and Reinforcement Learning','2026','Python, PyTorch, DreamerV3, PPO, adaptive computation, experiment diagnostics',[('Benchmarked adaptive halting and reinforcement-learning workflows, documenting sparse-retrieval failure modes, training tradeoffs, and reproducible evaluation results.','Benchmarked adaptive halting and reinforcement-learning workflows')],True)
project(m,'Badminton AI Match Analysis','2026','Python, computer vision, player tracking, pose estimation, video processing',[('Developed a two-player match-analysis pipeline with stable tracking, pose overlays, shuttle recovery, and annotated video visualization.','Developed a two-player match-analysis pipeline')],True)
_master_clear=project(m,'ClearMove - AI-Powered Rental Evidence Platform','Oct 2025 - Present','Next.js, TypeScript, PostgreSQL, Drizzle ORM, Gemini API, pdf-lib',[('Led 50+ stakeholder interviews and shipped a deployed evidence workflow across 21 API routes and 10 PostgreSQL tables, transforming lease and inspection inputs into auditable, timestamped PDFs.','Led 50+ stakeholder interviews')],True)
project(m,'TableUs - AI-Assisted Group Restaurant Discovery','2026','Next.js, React, Python, FastAPI, Gemini API, Google Maps API',[('Built an explainable group recommendation workflow combining user preferences, location, and natural-language input; earned Second Place at the 2026 Cursor Hackathon through an explainable, location-aware group decision flow.','Built an explainable group recommendation workflow')],True)
project(m,'Brain-Computer Interface MVP','2026','Python, Streamlit, real-time signal pipelines, interactive visualization',[('Built an end-to-end neural-signal prototype to process, inspect, and visualize streaming data for applied BCI experimentation.','Built an end-to-end neural-signal prototype')],True)
project(m,'Dungeons & Dragons Voice Adventure','2026','Gemini Live, multimodal voice interaction, dynamic visuals',[('Delivered a real-time storytelling system with turn-based gameplay, live agent responses, multimodal voice interaction, and dynamically generated visual feedback.','Delivered a real-time storytelling system')],True)
project(m,'prompttrace - LLM Observability','2026','LLM tracing, latency logging, failure analysis, cost analytics',[('Built experiment observability for LLM workflows, exposing failures, latency, and spend so model behavior could be debugged and optimized across repeated runs.','Built experiment observability for LLM workflows')],True)
project(m,'llm-fit - Local Model Hardware Advisor','2026','Local LLM hardware fit estimation, model selection, command generation',[('Translated device constraints into compatible local-model choices and runnable commands, helping users evaluate deployment feasibility before downloading or serving a model.','Translated device constraints into compatible local-model choices and runnable commands')],True)
section(m,'Open Source Contributions',True)
bullet(m,'Recorded 213 merged pull requests, including 201 merged contributions to external repositories, as verified through the GitHub API on September 22, 2026.','Recorded 213 merged pull requests',9.05,True)
bullet(m,'Contributed through maintainer review and CI workflows to GitHub Docs, Docker CLI, Microsoft APM, NVIDIA NemoClaw, Meta DotSlash, Repomix, and other developer tools.','Contributed through maintainer review and CI workflows',9.05,True)
section(m,'Leadership and Activities',True)
bullet(m,'Badger Blockchain E-Board, Social Media and Communications | Summer 2026 - Present','Badger Blockchain E-Board, Social Media and Communications',9.05,True)
bullet(m,'Badminton E-Board | Spring 2025 - Present','Badminton E-Board',9.05,True)
bullet(m,'China and U.S. Engagement at GWU - Founder and Head of Logistics | Aug 2024 - Present','China and U.S. Engagement at GWU - Founder and Head of Logistics',9.05,True)
bullet(m,'China Student Development Think Tank - Leader and Head Scheduler | Aug 2023 - Present','China Student Development Think Tank - Leader and Head Scheduler',9.05,True)
bullet(m,'UW-Madison Data Science Club - promoted technical events and organized weekly Python, R, and Excel workshops.','UW-Madison Data Science Club',9.05,True)
bullet(m,'VEX Robotics Club - built and maintained a C++ robot and evaluated routing strategies for competitive matches.','VEX Robotics Club',9.05,True)
achievements(m,['Second Place, 2026 Cursor Hackathon - TableUs','Third Prize, 2025 Badger Build Fest Venture Track - ClearMove','Second Place, 2024 GWU Data Science Association Datathon (Python)','Second Place (Skills) and Third Place (Qualification), 2023 VEX Robotics World Championship (C++)'],True)
master_path=save(m,'masters','Ching-Wei-Kang-Masters-Application-CV.docx')
print('\n'.join(map(str,[master_path,swe_path,ai_path,ds_path])))
